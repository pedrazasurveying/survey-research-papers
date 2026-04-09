"""
Subagent 5B: TABLES_P2
========================
Read Paper 2 analysis CSV outputs (demographic disparity and earnings gaps)
and produce APA 7 formatted publication-ready tables as .docx files.

Tables produced:
    1. Table 1 -- Demographic Composition of Surveyors and Survey Technicians
    2. Table 2 -- Representation of Surveyors Relative to U.S. Workforce
    3. Table 3 -- Median Annual Earnings by Demographic Group (Surveyors)
    4. Table 4 -- OLS Regression Results: Log Hourly Wage
    5. Table 5 -- Geographic Variation in Surveyor Demographics (Top 10 States)

Usage:
    python scripts/s5b_tables_p2.py
    python scripts/s5b_tables_p2.py --indir analysis/p2 --outdir tables/p2
"""

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd

from docx import Document
from docx.shared import Inches, Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ---------------------------------------------------------------------------
# APA 7 formatting constants
# ---------------------------------------------------------------------------
FONT_NAME = "Times New Roman"
FONT_SIZE_BODY = Pt(11)
FONT_SIZE_TITLE = Pt(12)
FONT_SIZE_NOTE = Pt(10)
TABLE_WIDTH = Inches(6.5)

# Suppression / em-dash character
EM_DASH = "\u2014"


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------
def fmt_pct(val, decimals=1):
    """Format percentage with specified decimals. Return em dash for NaN."""
    if pd.isna(val):
        return EM_DASH
    return f"{val:.{decimals}f}"


def fmt_dollars(val, decimals=0):
    """Format dollar value with comma separator. Return em dash for NaN."""
    if pd.isna(val):
        return EM_DASH
    if decimals == 0:
        return "$" + f"{val:,.0f}"
    return "$" + f"{val:,.{decimals}f}"


def fmt_int(val):
    """Format integer with comma separator. Return em dash for NaN."""
    if pd.isna(val):
        return EM_DASH
    return f"{int(val):,}"


def fmt_pval(p):
    """Format p-value with significance stars."""
    if pd.isna(p):
        return EM_DASH
    if p < 0.001:
        return "< .001***"
    elif p < 0.01:
        return f"{p:.3f}**"
    elif p < 0.05:
        return f"{p:.3f}*"
    else:
        return f"{p:.3f}"


def sig_stars(p):
    """Return significance stars only."""
    if pd.isna(p):
        return ""
    if p < 0.001:
        return "***"
    elif p < 0.01:
        return "**"
    elif p < 0.05:
        return "*"
    return ""


def fmt_coef(val, decimals=4):
    """Format a regression coefficient."""
    if pd.isna(val):
        return EM_DASH
    return f"{val:.{decimals}f}"


def fmt_pct_se(pct, se):
    """Format 'XX.X (Y.Y)' for percentage with SE."""
    if pd.isna(pct):
        return EM_DASH
    pct_str = f"{pct:.1f}"
    if pd.isna(se) or se == 0:
        return pct_str
    return f"{pct_str} ({se:.1f})"


def fmt_ci_dollars(median_val, ci_lower, ci_upper):
    """Format '$Median [$CI_lower, $CI_upper]'."""
    if pd.isna(median_val):
        return EM_DASH
    med = "$" + f"{median_val:,.0f}"
    lo = ("$" + f"{ci_lower:,.0f}") if not pd.isna(ci_lower) else "?"
    hi = ("$" + f"{ci_upper:,.0f}") if not pd.isna(ci_upper) else "?"
    return f"{med} [{lo}, {hi}]"


# ---------------------------------------------------------------------------
# APA 7 Document and Table construction helpers
# ---------------------------------------------------------------------------
def create_apa_document():
    """Create a new Document with APA 7 default styles."""
    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = FONT_NAME
    font.size = FONT_SIZE_BODY
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    return doc


def add_table_title(doc, title_text):
    """Add an APA 7 table title: bold, 12pt, flush left."""
    para = doc.add_paragraph()
    run = para.add_run(title_text)
    run.bold = True
    run.font.size = FONT_SIZE_TITLE
    run.font.name = FONT_NAME
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.space_before = Pt(12)
    return para


def add_table_note(doc, note_text):
    """Add an APA 7 table note in 10pt italic."""
    para = doc.add_paragraph()
    run = para.add_run("Note. ")
    run.italic = True
    run.font.size = FONT_SIZE_NOTE
    run.font.name = FONT_NAME
    run2 = para.add_run(note_text)
    run2.italic = True
    run2.font.size = FONT_SIZE_NOTE
    run2.font.name = FONT_NAME
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(12)
    return para


def set_cell_text(cell, text, bold=False, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """Set cell text with APA formatting."""
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(str(text))
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE_BODY
    run.bold = bold
    para.alignment = alignment
    # Reduce cell padding
    cell_xml = cell._tc
    tcPr = cell_xml.get_or_add_tcPr()
    tcMar = parse_xml(
        '<w:tcMar ' + nsdecls("w") + '>'
        '<w:top w:w="40" w:type="dxa"/>'
        '<w:bottom w:w="40" w:type="dxa"/>'
        '<w:left w:w="60" w:type="dxa"/>'
        '<w:right w:w="60" w:type="dxa"/>'
        '</w:tcMar>'
    )
    tcPr.append(tcMar)


def remove_all_borders(table):
    """Remove all cell borders from a table."""
    tbl = table._tbl
    for cell in tbl.iter_tcs():
        tcPr = cell.get_or_add_tcPr()
        tcBorders = parse_xml(
            '<w:tcBorders ' + nsdecls("w") + '>'
            '<w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
            '</w:tcBorders>'
        )
        tcPr.append(tcBorders)


def set_row_border_bottom(row, sz="4", color="000000"):
    """Add a bottom border to every cell in a row."""
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            '<w:tcBorders ' + nsdecls("w") + '>'
            '<w:bottom w:val="single" w:sz="' + sz + '" w:space="0" w:color="' + color + '"/>'
            '</w:tcBorders>'
        )
        tcPr.append(tcBorders)


def set_row_border_top(row, sz="4", color="000000"):
    """Add a top border to every cell in a row."""
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            '<w:tcBorders ' + nsdecls("w") + '>'
            '<w:top w:val="single" w:sz="' + sz + '" w:space="0" w:color="' + color + '"/>'
            '</w:tcBorders>'
        )
        tcPr.append(tcBorders)


def apply_apa_borders(table, header_row_count=1):
    """
    Apply APA 7 borders: top of table, below header, bottom of table.
    No vertical lines, no other horizontal lines.
    """
    remove_all_borders(table)
    set_row_border_top(table.rows[0], sz="8")
    set_row_border_bottom(table.rows[header_row_count - 1], sz="4")
    set_row_border_bottom(table.rows[-1], sz="8")


def load_csv_safe(path):
    """Load a CSV, returning None and printing a warning if the file is missing."""
    if not path.exists():
        print(f"  WARNING: {path} not found. Skipping dependent table.")
        return None
    df = pd.read_csv(path)
    print(f"  Loaded {path.name} ({len(df)} rows)")
    return df


# ---------------------------------------------------------------------------
# Table 1: Demographic Composition of Surveyors and Survey Technicians
# ---------------------------------------------------------------------------
def build_table1(doc, comp_df, chi2_df):
    """
    Panel A: Race/Ethnicity, Panel B: Sex, Panel C: Education, Panel D: Veteran Status.
    Columns: Category | Surveyors % (SE) | Technicians % (SE) | chi-square p-value
    """
    print("\n  Building Table 1: Demographic Composition...")

    add_table_title(
        doc,
        "Table 1\n"
        "Demographic Composition of Surveyors and Survey Technicians, "
        "ACS 2020\u20132024"
    )

    # Panel definitions: (panel_label, variable_name, category_order)
    panels = [
        ("Panel A: Race/Ethnicity", "race_eth",
         ["NH White", "NH Black", "Hispanic", "NH Asian", "NH Other"]),
        ("Panel B: Sex", "sex_r",
         ["Male", "Female"]),
        ("Panel C: Education", "educ_r",
         ["HS or less", "Some college or AA", "BA or BS", "Graduate degree"]),
        ("Panel D: Veteran Status", "veteran",
         ["Veteran", "Non-veteran"]),
    ]

    # Count total data rows + panel headers + column header
    total_rows = 1  # header row
    for panel_label, var, cats in panels:
        total_rows += 1  # panel label row
        cats_in_data = [c for c in cats
                        if len(comp_df[(comp_df["variable"] == var)
                                       & (comp_df["category"] == c)]) > 0]
        total_rows += len(cats_in_data)

    num_cols = 4
    table = doc.add_table(rows=total_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header row
    headers = ["Demographic Group", "Surveyors % (SE)", "Technicians % (SE)",
               "\u03c7\u00b2 p-value"]
    for j, h in enumerate(headers):
        align = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(table.rows[0].cells[j], h, bold=True, alignment=align)

    row_idx = 1
    for panel_label, var, cat_order in panels:
        # Panel label row (merge all cells)
        row = table.rows[row_idx]
        for ci in range(1, num_cols):
            row.cells[0].merge(row.cells[ci])
        set_cell_text(row.cells[0], panel_label, bold=True)
        row_idx += 1

        # Get chi2 p-value for this variable
        chi2_p = None
        if chi2_df is not None and len(chi2_df) > 0:
            chi2_row = chi2_df[chi2_df["variable"] == var]
            if len(chi2_row) > 0:
                chi2_p = chi2_row.iloc[0]["p_value"]

        cats_in_data = [c for c in cat_order
                        if len(comp_df[(comp_df["variable"] == var)
                                       & (comp_df["category"] == c)]) > 0]

        for k, cat in enumerate(cats_in_data):
            row = table.rows[row_idx]

            # Category label (indented)
            set_cell_text(row.cells[0], f"  {cat}")

            # Surveyors
            sv = comp_df[(comp_df["variable"] == var) &
                         (comp_df["category"] == cat) &
                         (comp_df["occupation"] == "Surveyors")]
            if len(sv) > 0:
                sv = sv.iloc[0]
                set_cell_text(row.cells[1],
                              fmt_pct_se(sv["weighted_pct"], sv["se"]),
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                set_cell_text(row.cells[1], EM_DASH,
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)

            # Technicians
            tc = comp_df[(comp_df["variable"] == var) &
                         (comp_df["category"] == cat) &
                         (comp_df["occupation"] == "Technicians")]
            if len(tc) > 0:
                tc = tc.iloc[0]
                set_cell_text(row.cells[2],
                              fmt_pct_se(tc["weighted_pct"], tc["se"]),
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                set_cell_text(row.cells[2], EM_DASH,
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)

            # chi-square p-value (only on first category row of each panel)
            if k == 0 and chi2_p is not None:
                set_cell_text(row.cells[3], fmt_pval(chi2_p),
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                set_cell_text(row.cells[3], "",
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)

            row_idx += 1

    apply_apa_borders(table, header_row_count=1)

    add_table_note(
        doc,
        "Weighted percentages with BRR standard errors (Fay k = 0.5) "
        "in parentheses. \u03c7\u00b2 test compares distributions between "
        "Surveyors (OCC 1310) and Technicians (OCC 1560). "
        "Data source: IPUMS ACS 2020\u20132024 5-year pooled microdata. "
        "***p < .001, **p < .01, *p < .05."
    )

    print("    Table 1 complete.")
    return True


# ---------------------------------------------------------------------------
# Table 2: Representation of Surveyors Relative to U.S. Workforce
# ---------------------------------------------------------------------------
def build_table2(doc, ratios_df):
    """
    Columns: Demographic Group | Surveyor Share (%) | U.S. Workforce Share (%) |
             Representation Ratio | Assessment
    """
    print("\n  Building Table 2: Representation Ratios...")

    add_table_title(
        doc,
        "Table 2\n"
        "Representation of Surveyors Relative to the U.S. Civilian Workforce, "
        "ACS 2020\u20132024"
    )

    # Category display order
    cat_order = ["Male", "Female", "NH White", "NH Black", "Hispanic",
                 "NH Asian", "Veteran"]
    display_df = ratios_df.set_index("category").reindex(
        [c for c in cat_order if c in ratios_df["category"].values]
    ).reset_index()

    num_cols = 5
    num_rows = 1 + len(display_df)  # header + data
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header
    headers = ["Demographic Group", "Surveyor Share (%)",
               "U.S. Workforce Share (%)", "Representation Ratio", "Assessment"]
    for j, h in enumerate(headers):
        align = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(table.rows[0].cells[j], h, bold=True, alignment=align)

    for i, (_, row_data) in enumerate(display_df.iterrows()):
        row = table.rows[i + 1]
        set_cell_text(row.cells[0], row_data["category"])
        set_cell_text(row.cells[1], fmt_pct(row_data["surveyor_pct"]),
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[2], fmt_pct(row_data["us_workforce_pct"]),
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

        ratio = row_data["representation_ratio"]
        ratio_str = f"{ratio:.2f}" if not pd.isna(ratio) else EM_DASH
        set_cell_text(row.cells[3], ratio_str,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

        flag = row_data.get("flag", "")
        if pd.isna(flag):
            flag = ""
        assessment = flag if flag else "Proportionate"
        if "SEVERE" in str(flag).upper():
            assessment = "Severe under-repr."
        elif "under" in str(flag).lower():
            assessment = "Under-represented"
        set_cell_text(row.cells[4], assessment,
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)

    apply_apa_borders(table, header_row_count=1)

    add_table_note(
        doc,
        "Representation ratio = surveyor share / U.S. workforce share. "
        "Ratios below 0.80 indicate under-representation; below 0.50 "
        "indicates severe under-representation. "
        "Data source: IPUMS ACS 2020\u20132024."
    )

    print("    Table 2 complete.")
    return True


# ---------------------------------------------------------------------------
# Table 3: Median Annual Earnings by Demographic Group (Surveyors)
# ---------------------------------------------------------------------------
def build_table3(doc, earnings_df, gaps_df):
    """
    Gender rows, Race/ethnicity rows, Veteran rows.
    Columns: Group | Median Earnings (95% CI) | n | Gap vs. Reference ($) | Gap (%)
    Reference groups: Male (for gender), NH White (for race).
    """
    print("\n  Building Table 3: Earnings by Demographic Group...")

    add_table_title(
        doc,
        "Table 3\n"
        "Median Annual Earnings by Demographic Group Among Surveyors, "
        "ACS 2020\u20132024"
    )

    # Build row data from wage_gaps and earnings_core
    # Filter to Surveyors, annual wage measure only
    sv_gaps = None
    if gaps_df is not None:
        sv_gaps = gaps_df[
            (gaps_df["occupation"] == "Surveyors") &
            (gaps_df["measure"] == "wage_adj")
        ].copy()

    # Build table rows: (label, median, ci_lower, ci_upper, n, gap_dollars, gap_pct, suppressed)
    table_rows = []

    # --- Gender section ---
    table_rows.append(("Gender", None, None, None, None, None, None, True))

    if sv_gaps is not None:
        gender_row = sv_gaps[sv_gaps["comparison"].str.contains("Gender")]
        if len(gender_row) > 0:
            gr = gender_row.iloc[0]
            # Male (reference)
            if not pd.isna(gr["group_a_median"]):
                ci_half = 1.96 * gr["group_a_se"] if not pd.isna(gr["group_a_se"]) else 0
                table_rows.append((
                    "  Male (ref.)",
                    gr["group_a_median"],
                    gr["group_a_median"] - ci_half,
                    gr["group_a_median"] + ci_half,
                    gr["group_a_n"],
                    EM_DASH,
                    EM_DASH,
                    False,
                ))
            # Female
            if gr.get("suppressed", False):
                table_rows.append(("  Female", None, None, None, None, None, None, True))
            else:
                ci_half_b = 1.96 * gr["group_b_se"] if not pd.isna(gr["group_b_se"]) else 0
                table_rows.append((
                    "  Female",
                    gr["group_b_median"],
                    gr["group_b_median"] - ci_half_b,
                    gr["group_b_median"] + ci_half_b,
                    gr["group_b_n"],
                    gr["gap_dollars"],
                    gr["gap_pct"],
                    False,
                ))

    # --- Race/ethnicity section ---
    table_rows.append(("Race/Ethnicity", None, None, None, None, None, None, True))

    race_order = ["NH White", "NH Black", "Hispanic", "NH Asian", "NH Other"]
    if sv_gaps is not None:
        # NH White reference
        white_row = sv_gaps[
            sv_gaps["comparison"].str.contains("Race") &
            (sv_gaps["group_a"] == "NH White")
        ]
        if len(white_row) > 0:
            wr = white_row.iloc[0]
            ci_half = 1.96 * wr["group_a_se"] if not pd.isna(wr["group_a_se"]) else 0
            table_rows.append((
                "  NH White (ref.)",
                wr["group_a_median"],
                wr["group_a_median"] - ci_half,
                wr["group_a_median"] + ci_half,
                wr["group_a_n"],
                EM_DASH,
                EM_DASH,
                False,
            ))

        for race in race_order:
            if race == "NH White":
                continue
            race_gaps = sv_gaps[
                sv_gaps["comparison"].str.contains("Race") &
                (sv_gaps["group_b"] == race)
            ]
            if len(race_gaps) > 0:
                rg = race_gaps.iloc[0]
                if rg.get("suppressed", False):
                    table_rows.append((f"  {race}", None, None, None,
                                       rg["group_b_n"], None, None, True))
                else:
                    ci_half_b = 1.96 * rg["group_b_se"] if not pd.isna(rg["group_b_se"]) else 0
                    table_rows.append((
                        f"  {race}",
                        rg["group_b_median"],
                        rg["group_b_median"] - ci_half_b,
                        rg["group_b_median"] + ci_half_b,
                        rg["group_b_n"],
                        rg["gap_dollars"],
                        rg["gap_pct"],
                        False,
                    ))

    # --- Veteran section ---
    table_rows.append(("Veteran Status", None, None, None, None, None, None, True))

    if sv_gaps is not None:
        for vet_label in ["Non-veteran", "Veteran"]:
            vet_rows = sv_gaps[
                sv_gaps["comparison"].str.contains("Veteran") &
                (sv_gaps["group_a"] == vet_label)
            ]
            if len(vet_rows) > 0:
                vr = vet_rows.iloc[0]
                if not pd.isna(vr["group_a_median"]):
                    ci_half = 1.96 * vr["group_a_se"] if not pd.isna(vr["group_a_se"]) else 0
                    table_rows.append((
                        f"  {vet_label}",
                        vr["group_a_median"],
                        vr["group_a_median"] - ci_half,
                        vr["group_a_median"] + ci_half,
                        vr["group_a_n"],
                        EM_DASH,
                        EM_DASH,
                        False,
                    ))

    # Create the docx table
    num_cols = 5
    num_rows = 1 + len(table_rows)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header
    col_headers = ["Group", "Median Earnings (95% CI)", "n",
                   "Gap vs. Ref. ($)", "Gap (%)"]
    for j, h in enumerate(col_headers):
        align = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(table.rows[0].cells[j], h, bold=True, alignment=align)

    for i, (label, median, ci_lo, ci_hi, n, gap_d, gap_p, is_header) in enumerate(table_rows):
        row = table.rows[i + 1]
        if is_header and median is None:
            # Section header -- merge and bold
            for ci in range(1, num_cols):
                row.cells[0].merge(row.cells[ci])
            set_cell_text(row.cells[0], label, bold=True)
        else:
            set_cell_text(row.cells[0], label)
            set_cell_text(row.cells[1], fmt_ci_dollars(median, ci_lo, ci_hi),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            n_str = EM_DASH
            if n is not None:
                if isinstance(n, (int, float)) and not pd.isna(n):
                    n_str = fmt_int(n)
                elif isinstance(n, str):
                    n_str = n
            set_cell_text(row.cells[2], n_str,
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)

            # Gap columns
            if isinstance(gap_d, str):
                set_cell_text(row.cells[3], gap_d,
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                set_cell_text(row.cells[3], fmt_dollars(gap_d),
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)

            if isinstance(gap_p, str):
                set_cell_text(row.cells[4], gap_p,
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                gap_p_str = fmt_pct(gap_p) if not pd.isna(gap_p) else EM_DASH
                set_cell_text(row.cells[4], gap_p_str,
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)

    apply_apa_borders(table, header_row_count=1)

    add_table_note(
        doc,
        "Weighted median annual earnings (inflation-adjusted to 2024 $) with "
        "95% confidence intervals from BRR standard errors (Woodruff method). "
        "Gap computed as reference group median minus comparison group median. "
        "Cells with unweighted n < 50 are suppressed (\u2014). "
        "Data source: IPUMS ACS 2020\u20132024."
    )

    print("    Table 3 complete.")
    return True


# ---------------------------------------------------------------------------
# Table 4: OLS Regression Results -- Log Hourly Wage
# ---------------------------------------------------------------------------
def build_table4(doc, reg_df):
    """
    Standard regression table: Variable | B | SE | t | p | Interpretation
    Note about state fixed effects, HC1 robust SE.
    """
    print("\n  Building Table 4: OLS Regression Results...")

    add_table_title(
        doc,
        "Table 4\n"
        "OLS Regression Results: Log Hourly Wage Among Surveyors, "
        "ACS 2020\u20132024"
    )

    # Separate model summary from coefficients
    summary_row = reg_df[reg_df["variable"] == "MODEL_SUMMARY"]
    coef_df = reg_df[reg_df["variable"] != "MODEL_SUMMARY"].copy()

    # Filter to key variable groups (exclude state FE rows for display)
    display_groups = ["other", "race_eth", "sex_r", "educ_r", "agegroup",
                      "veteran", "classwkr"]
    display_df = coef_df[coef_df["group"].isin(display_groups)].copy()

    # Friendly variable names
    name_map = {
        "const": "Intercept",
    }
    # Build dynamic name map from actual variable names
    for _, row in display_df.iterrows():
        var = row["variable"]
        if var.startswith("race_eth_"):
            name_map[var] = var.replace("race_eth_", "Race: ")
        elif var.startswith("sex_r_"):
            name_map[var] = var.replace("sex_r_", "Sex: ")
        elif var.startswith("educ_r_"):
            name_map[var] = var.replace("educ_r_", "Education: ")
        elif var.startswith("agegroup_"):
            name_map[var] = var.replace("agegroup_", "Age: ")
        elif var.startswith("veteran_"):
            name_map[var] = var.replace("veteran_", "Veteran: ")
        elif var.startswith("classwkr_"):
            name_map[var] = var.replace("classwkr_", "Class of worker: ")

    # Variable group ordering
    group_order = ["other", "sex_r", "race_eth", "educ_r", "agegroup",
                   "veteran", "classwkr"]
    group_labels = {
        "other": None,
        "sex_r": "Sex (ref: Male)",
        "race_eth": "Race/Ethnicity (ref: NH White)",
        "educ_r": "Education (ref: HS or less)",
        "agegroup": "Age Group (ref: 25\u201334)",
        "veteran": "Veteran Status (ref: Non-veteran)",
        "classwkr": "Class of Worker (ref: Private wage/salary)",
    }

    # Build ordered list of display rows
    ordered_rows = []
    for grp in group_order:
        grp_rows = display_df[display_df["group"] == grp]
        if len(grp_rows) == 0:
            continue
        label = group_labels.get(grp)
        if label:
            ordered_rows.append(("SECTION", label, None))
        for _, r in grp_rows.iterrows():
            display_name = name_map.get(r["variable"], r["variable"])
            ordered_rows.append(("DATA", display_name, r))

    # Add state FE indicator row
    n_state_fe = len(coef_df[coef_df["group"] == "state_fe"])

    num_cols = 6
    num_data_rows = len(ordered_rows) + 1 + 1  # +1 header, +1 state FE indicator
    # Add model summary rows
    has_summary = len(summary_row) > 0
    if has_summary:
        num_data_rows += 3  # R2, Adj R2, N

    table = doc.add_table(rows=num_data_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header
    headers = ["Variable", "B", "SE", "t", "p", "% Wage Effect"]
    for j, h in enumerate(headers):
        align = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(table.rows[0].cells[j], h, bold=True, alignment=align)

    row_idx = 1
    for row_type, label, data in ordered_rows:
        row = table.rows[row_idx]
        if row_type == "SECTION":
            for ci in range(1, num_cols):
                row.cells[0].merge(row.cells[ci])
            set_cell_text(row.cells[0], label, bold=True)
        else:
            set_cell_text(row.cells[0], f"  {label}")
            set_cell_text(row.cells[1], fmt_coef(data["coefficient"]),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[2], fmt_coef(data["se"]),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            t_val = data.get("t_stat", None)
            set_cell_text(row.cells[3], fmt_coef(t_val, 2),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)

            p_val = data["p_value"]
            stars = sig_stars(p_val)
            p_str = fmt_coef(p_val, 3) + stars if not pd.isna(p_val) else EM_DASH
            set_cell_text(row.cells[4], p_str,
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)

            # Percentage wage effect = (exp(B) - 1) * 100
            coeff = data["coefficient"]
            if not pd.isna(coeff) and data["variable"] != "const":
                pct_effect = (np.exp(coeff) - 1) * 100
                set_cell_text(row.cells[5],
                              f"{pct_effect:+.1f}%",
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                set_cell_text(row.cells[5], EM_DASH,
                              alignment=WD_ALIGN_PARAGRAPH.CENTER)
        row_idx += 1

    # State FE indicator row
    row = table.rows[row_idx]
    set_cell_text(row.cells[0], "State fixed effects")
    for ci in range(1, num_cols - 1):
        set_cell_text(row.cells[ci], "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    indicator = f"Yes ({n_state_fe})" if n_state_fe > 0 else "No"
    set_cell_text(row.cells[num_cols - 1], indicator,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER)
    row_idx += 1

    # Model summary rows
    if has_summary:
        sm = summary_row.iloc[0]
        for stat_label, stat_val, fmt_func in [
            ("R\u00b2", sm.get("r_squared"), lambda v: f"{v:.4f}"),
            ("Adjusted R\u00b2", sm.get("adj_r_squared"), lambda v: f"{v:.4f}"),
            ("N", sm.get("n_obs"), lambda v: f"{int(v):,}"),
        ]:
            row = table.rows[row_idx]
            set_cell_text(row.cells[0], stat_label, bold=True)
            for ci in range(1, num_cols - 1):
                set_cell_text(row.cells[ci], "", alignment=WD_ALIGN_PARAGRAPH.CENTER)
            val_str = fmt_func(stat_val) if not pd.isna(stat_val) else EM_DASH
            set_cell_text(row.cells[num_cols - 1], val_str,
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            row_idx += 1

    apply_apa_borders(table, header_row_count=1)

    add_table_note(
        doc,
        "Dependent variable: log(hourly wage). Estimated via OLS with "
        "heteroscedasticity-consistent (HC1) robust standard errors. "
        "Reference categories: Male, NH White, HS or less, age 25\u201334, "
        "Non-veteran, Private wage/salary. "
        "% Wage Effect = (exp(B) \u2212 1) \u00d7 100. "
        "***p < .001, **p < .01, *p < .05."
    )

    print("    Table 4 complete.")
    return True


# ---------------------------------------------------------------------------
# Table 5: Geographic Variation in Surveyor Demographics (Top 10 States)
# ---------------------------------------------------------------------------
def build_table5(doc, state_df, metro_df):
    """
    Columns: State | N | % Non-White | % Female | Median Wage
    Include metro vs non-metro summary row.
    """
    print("\n  Building Table 5: Geographic Variation...")

    add_table_title(
        doc,
        "Table 5\n"
        "Geographic Variation in Surveyor Demographics: Top 10 States "
        "by Workforce Size, ACS 2020\u20132024"
    )

    # Get top 10 states by weighted_n
    top10 = state_df.nlargest(10, "weighted_n").copy()

    # Build table: header + 10 states + separator + metro rows
    metro_rows_data = []
    if metro_df is not None and len(metro_df) > 0:
        for _, mr in metro_df.iterrows():
            metro_rows_data.append(mr)

    num_cols = 5
    num_rows = 1 + len(top10) + len(metro_rows_data)  # header + states + metro
    if len(metro_rows_data) > 0:
        num_rows += 1  # separator / section label row

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header
    headers = ["State", "N (weighted)", "% Non-White", "% Female", "Median Wage"]
    for j, h in enumerate(headers):
        align = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
        set_cell_text(table.rows[0].cells[j], h, bold=True, alignment=align)

    row_idx = 1
    for _, st in top10.iterrows():
        row = table.rows[row_idx]
        set_cell_text(row.cells[0], st["state_name"])
        set_cell_text(row.cells[1], fmt_int(st["weighted_n"]),
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[2], fmt_pct(st.get("pct_nonwhite", np.nan)),
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[3], fmt_pct(st.get("pct_female", np.nan)),
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[4], fmt_dollars(st.get("median_wage_adj", np.nan)),
                      alignment=WD_ALIGN_PARAGRAPH.CENTER)
        row_idx += 1

    # Metro / non-metro section
    if len(metro_rows_data) > 0:
        # Section label
        row = table.rows[row_idx]
        for ci in range(1, num_cols):
            row.cells[0].merge(row.cells[ci])
        set_cell_text(row.cells[0], "Metro Status", bold=True)
        set_row_border_bottom(table.rows[row_idx], sz="2", color="666666")
        row_idx += 1

        for mr in metro_rows_data:
            row = table.rows[row_idx]
            set_cell_text(row.cells[0], f"  {mr['metro_status']}")
            set_cell_text(row.cells[1], fmt_int(mr.get("weighted_n", np.nan)),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[2], fmt_pct(mr.get("pct_nonwhite", np.nan)),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[3], fmt_pct(mr.get("pct_female", np.nan)),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[4], fmt_dollars(mr.get("median_wage_adj", np.nan)),
                          alignment=WD_ALIGN_PARAGRAPH.CENTER)
            row_idx += 1

    apply_apa_borders(table, header_row_count=1)

    add_table_note(
        doc,
        "States ranked by estimated surveyor workforce size (weighted N). "
        "Median wages inflation-adjusted to 2024 $. "
        "States with unweighted n < 30 are excluded. "
        "Metro status classification follows IPUMS METRO variable. "
        "Data source: IPUMS ACS 2020\u20132024."
    )

    print("    Table 5 complete.")
    return True


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(
        description="Subagent 5B: Publication-ready APA 7 tables for Paper 2 "
                    "(Demographic Disparity and Earnings Gaps)"
    )
    parser.add_argument(
        "--indir", default="analysis/p2",
        help="Directory containing Paper 2 analysis CSVs (default: analysis/p2)",
    )
    parser.add_argument(
        "--outdir", default="tables/p2",
        help="Output directory for .docx tables (default: tables/p2)",
    )
    args = parser.parse_args()

    indir = Path(args.indir)
    outdir = Path(args.outdir)

    if not indir.exists():
        print(f"ERROR: Input directory does not exist: {indir}")
        print("Run s4b_disparity_analysis.py and s4c_earnings_analysis.py first.")
        sys.exit(1)

    outdir.mkdir(parents=True, exist_ok=True)

    print("=" * 65)
    print("PAPER 2: APA 7 TABLE GENERATION")
    print("=" * 65)
    print(f"  Input directory:  {indir}")
    print(f"  Output directory: {outdir}")

    # ------------------------------------------------------------------
    # Load all required CSVs
    # ------------------------------------------------------------------
    print("\nLoading analysis CSVs...")

    comp_df = load_csv_safe(indir / "p2_composition.csv")
    chi2_df = load_csv_safe(indir / "p2_composition_chi2.csv")
    ratios_df = load_csv_safe(indir / "p2_representation_ratios.csv")
    earnings_df = load_csv_safe(indir / "p2_earnings_core.csv")
    gaps_df = load_csv_safe(indir / "p2_wage_gaps.csv")
    reg_df = load_csv_safe(indir / "p2_wage_regression.csv")
    state_df = load_csv_safe(indir / "p2_state_composition.csv")
    metro_df = load_csv_safe(indir / "p2_metro_nonmetro.csv")

    # ------------------------------------------------------------------
    # Generate tables
    # ------------------------------------------------------------------
    n_produced = 0

    # Table 1: Demographic Composition
    if comp_df is not None:
        try:
            doc = create_apa_document()
            if build_table1(doc, comp_df, chi2_df):
                outpath = outdir / "table1_demographic_composition.docx"
                doc.save(str(outpath))
                print(f"    -> {outpath}")
                n_produced += 1
        except Exception as e:
            print(f"    ERROR building Table 1: {e}")

    # Table 2: Representation Ratios
    if ratios_df is not None:
        try:
            doc = create_apa_document()
            if build_table2(doc, ratios_df):
                outpath = outdir / "table2_representation_ratios.docx"
                doc.save(str(outpath))
                print(f"    -> {outpath}")
                n_produced += 1
        except Exception as e:
            print(f"    ERROR building Table 2: {e}")

    # Table 3: Earnings
    if gaps_df is not None:
        try:
            doc = create_apa_document()
            if build_table3(doc, earnings_df, gaps_df):
                outpath = outdir / "table3_earnings.docx"
                doc.save(str(outpath))
                print(f"    -> {outpath}")
                n_produced += 1
        except Exception as e:
            print(f"    ERROR building Table 3: {e}")

    # Table 4: Regression
    if reg_df is not None and len(reg_df) > 0:
        try:
            doc = create_apa_document()
            if build_table4(doc, reg_df):
                outpath = outdir / "table4_regression.docx"
                doc.save(str(outpath))
                print(f"    -> {outpath}")
                n_produced += 1
        except Exception as e:
            print(f"    ERROR building Table 4: {e}")

    # Table 5: Geographic
    if state_df is not None and len(state_df) > 0:
        try:
            doc = create_apa_document()
            if build_table5(doc, state_df, metro_df):
                outpath = outdir / "table5_geographic.docx"
                doc.save(str(outpath))
                print(f"    -> {outpath}")
                n_produced += 1
        except Exception as e:
            print(f"    ERROR building Table 5: {e}")

    # ------------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------------
    print("\n" + "=" * 65)
    print("PAPER 2 TABLE GENERATION COMPLETE")
    print("=" * 65)
    print(f"  {n_produced} table(s) produced in: {outdir}")
    if n_produced > 0:
        for f in sorted(outdir.glob("table*.docx")):
            size_kb = f.stat().st_size / 1024
            print(f"    {f.name}  ({size_kb:.1f} KB)")
    else:
        print("  No tables were produced. Ensure analysis CSVs exist in "
              f"{indir}.")
        print("  Run s4b_disparity_analysis.py, s4c_earnings_analysis.py, "
              "and s4d_geographic_analysis.py first.")


if __name__ == "__main__":
    main()
